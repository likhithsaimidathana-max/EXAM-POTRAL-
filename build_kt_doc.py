from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_ORIENT
from datetime import date

OUT = 'SmartExam_Pro_KT_Handover_Document.docx'
NAVY = '10233F'; GOLD = 'C89B3C'; TEAL = '1C6E78'; INK = '202630'; MUTED = '65758B'; PALE = 'F4F7FA'; LINE = 'D9E1EA'; RED = 'B42318'

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)
def border(cell, color=LINE):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for e in ('top','left','bottom','right'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),color); b.append(x)
    tcPr.append(b)
def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); m=tcPr.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); tcPr.append(m)
    for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=m.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); m.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)
def add_rule(p, color=GOLD):
    pPr=p._p.get_or_add_pPr(); pBdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'16'); bottom.set(qn('w:space'),'6'); bottom.set(qn('w:color'),color); pBdr.append(bottom); pPr.append(pBdr)
def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name='Aptos'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos')
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
    if italic is not None: run.italic=italic
def add_text(p, text, size=11, bold=False, color=INK, italic=False):
    r=p.add_run(text); set_font(r,size,bold,color,italic); return r
def style_doc(d):
    sec=d.sections[0]; sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    styles=d.styles
    for name,sz,col,bold in [('Normal',10.3,INK,False),('Title',28,NAVY,True),('Subtitle',13,MUTED,False),('Heading 1',17,NAVY,True),('Heading 2',12.5,TEAL,True),('Heading 3',11,GOLD,True)]:
        s=styles[name]; s.font.name='Aptos'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); s.font.size=Pt(sz); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=bold
    styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.12
    for n in ('Heading 1','Heading 2','Heading 3'):
        styles[n].paragraph_format.space_before=Pt(13); styles[n].paragraph_format.space_after=Pt(5)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.LEFT; add_text(header,'SMARTEXAM PRO  /  TECHNICAL HANDOVER',8,True,MUTED); add_rule(header, GOLD)
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_text(foot,'Confidential internal handover  |  SmartExam Pro',8,False,MUTED)
def para(d, text='', style=None, color=INK, size=10.3, bold=False):
    p=d.add_paragraph(style=style); add_text(p,text,size,bold,color); return p
def bullet(d,text):
    p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); add_text(p,text,10.2,False,INK); return p
def table(d, headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); border(c,NAVY); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; add_text(c.paragraphs[0],h,9.2,True,'FFFFFF')
        if widths: c.width=Inches(widths[i])
    set_repeat_table_header(t.rows[0])
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; set_cell_margins(c); border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx%2==1: shade(c,PALE)
            add_text(c.paragraphs[0],str(val),9.1,False,INK)
            if widths: c.width=Inches(widths[i])
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def callout(d,title,text,accent=GOLD):
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,'FFF8EC' if accent==GOLD else 'EEF7F8'); border(c,accent); set_cell_margins(c,140,180,140,180)
    p=c.paragraphs[0]; add_text(p,title+'  ',10,True,accent); add_text(p,text,9.8,False,INK)
    d.add_paragraph().paragraph_format.space_after=Pt(2)

d=Document(); style_doc(d)
# cover
d.add_paragraph().paragraph_format.space_after=Pt(36)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; add_text(p,'ENGINEERING HANDOVER',10,True,GOLD); add_rule(p,GOLD)
p=d.add_paragraph(style='Title'); add_text(p,'SmartExam Pro',30,True,NAVY)
p=d.add_paragraph(style='Subtitle'); add_text(p,'Full Project Knowledge Transfer Document',14,False,MUTED)
d.add_paragraph().paragraph_format.space_after=Pt(24)
t=d.add_table(rows=4,cols=2); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style='Table Grid'
for i,(a,b) in enumerate([('Application','Online examination management platform'),('Architecture','PHP + MySQL monolith, served via XAMPP/Apache'),('Repository','C:\\xampp\\htdocs\\exam'),('Prepared','18 August 2026')]):
    for j,x in enumerate((a,b)):
        c=t.cell(i,j); set_cell_margins(c,120,150,120,150); border(c); shade(c,PALE if j==0 else 'FFFFFF'); add_text(c.paragraphs[0],x,10, j==0, NAVY if j==0 else INK)
d.add_paragraph().paragraph_format.space_after=Pt(80)
callout(d,'Purpose','This document equips a new owner to run, support, troubleshoot, and safely evolve SmartExam Pro. Findings are grounded in the repository as reviewed on 18 August 2026.',GOLD)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_text(p,'Prepared for project continuity',9,False,MUTED,True)
d.add_page_break()

para(d,'Executive briefing','Heading 1')
para(d,'SmartExam Pro is a browser-based examination platform for issuing role-specific assessments, capturing answers, and producing immediate scores. It is implemented as server-rendered PHP pages backed by MySQL and runs locally under XAMPP. The solution is functional but has no framework, migrations, dependency manifest, or centralized authorization layer; ownership should prioritize stabilization before feature expansion.')
table(d,['Area','What the system delivers','Primary owner'],[
['Student experience','Registration, login, upcoming exams, passkey entry, timed MCQ attempt, score history and PDF output.','Application support'],
['Admin experience','Exam and section creation, in-browser question drafting, candidate/passkey selection, score review.','Exam operations'],
['Data & grading','MySQL persistence for users, exam definition, question bank, assigned passkeys, answers and aggregate scores.','Database owner'],
['Runtime','Apache/PHP through XAMPP; MySQL database named exam; CDN-hosted Bootstrap and SweetAlert2.','Infrastructure support'],
], [1.35,4.55,1.3])
para(d,'Critical handover note','Heading 2')
callout(d,'Release readiness','The repository references save_questions.php from add_question.php, but that file is not present in the project root. Exam-question publishing will fail until the missing endpoint is restored or the client flow is changed.',RED)
para(d,'Scope and assumptions','Heading 2')
bullet(d,'The documentation reflects source files available in C:\\xampp\\htdocs\\exam; no live database dump or deployed environment was inspected.')
bullet(d,'Database structures are derived from queries and the included schema fixer. Column types, indexes and constraints must be validated against the target MySQL instance.')
bullet(d,'The system uses application pages as endpoints; it does not expose a versioned REST API.')

para(d,'Solution architecture','Heading 1')
para(d,'The application is a classic PHP monolith: each page contains its own presentation, request handling and database access. config.php creates a global mysqli connection; PHP sessions retain authenticated user state; the browser calls submitexam.php and the missing save_questions.php through fetch().')
table(d,['Layer','Implementation','Notes'],[
['Presentation','Server-rendered HTML, Bootstrap 4/5 CDN, custom CSS, SweetAlert2','Mixed Bootstrap versions are used across pages.'],
['Application','Standalone PHP pages with inline SQL and JavaScript','No routing, service, repository or middleware abstraction.'],
['Data','MySQL / mysqli','Local defaults: localhost, root, blank password, database exam.'],
['Session','PHP native sessions','User flags: user_logged_in, registration_id, fullname; admin_logged_in for admin.'],
['Reporting','Bundled FPDF','Used by f_pdf.php for results export.'],
], [1.2,2.3,3.7])
para(d,'Runtime request path','Heading 2')
table(d,['Step','Flow'],[
['1','Browser requests a PHP page through Apache at http://localhost/exam/.'],
['2','Page starts/reuses PHP session and includes config.php where needed.'],
['3','Page executes mysqli queries directly against the exam database.'],
['4','HTML is rendered; JavaScript fetch calls are used for async submission/publishing.'],
['5','Responses redirect to dashboards or return JSON for UI alerts.'],
], [0.8,6.4])

para(d,'Repository map','Heading 1')
table(d,['Area','Key files','Responsibility'],[
['Configuration','config.php','Creates MySQL connection and selects exam database.'],
['Authentication','login.php, registration.php, edit_profile.php','Student registration and authentication; hardcoded admin login.'],
['Admin','admin.php, dashboard.php, create_exam.php, add_question.php, admin_student_scores.php','Exam authoring, question drafting, score reporting.'],
['Exam delivery','upcoming_exam.php, enter_exam_passkey.php, exam_instraction.php, start_exam.php, startexam.php, submitexam.php','Eligibility, passkey validation, attempt UI and server-side scoring.'],
['Student history','user.php, pastexam.php, f_pdf.php','Dashboard, results and PDF result export.'],
['Database utility','fix_db_schema.php','Inspects questions and conditionally adds section_id FK. Do not expose publicly.'],
['Assets / libraries','fpdf.php, font/, doc/, tutorial/, *.css, icon box.png','Bundled PDF library, documentation, fonts and UI assets.'],
['Tests','run-tests.ps1, tests/','PowerShell runner exists, but the expected tests\\run-tests.php was not found during review.'],
], [1.3,3.0,2.9])

para(d,'Business workflows','Heading 1')
para(d,'Admin: create and publish an assessment','Heading 2')
table(d,['Step','Page / component','Outcome'],[
['1','login.php','Admin authenticates with the hardcoded email/password check.'],
['2','create_exam.php','Creates exam record and one or more sections. CSRF token is validated.'],
['3','add_question.php','Selects eligible users by skill, drafts MCQ questions against section limits, generates passkeys in the browser.'],
['4','save_questions.php (referenced, missing)','Expected to persist questions and assigned passkeys. This is the current release blocker.'],
['5','admin_student_scores.php','Views submitted student scores and percentage.'],
], [0.55,2.1,4.55])
para(d,'Student: take an assessment','Heading 2')
table(d,['Step','Page / component','Outcome'],[
['1','registration.php / login.php','Student creates account and signs in; password_verify() checks stored hash.'],
['2','user.php → upcoming_exam.php','Student sees eligible exams for today/tomorrow. Grace window controls start availability.'],
['3','enter_exam_passkey.php','One-time passkey is validated for the student/exam pair and marked used.'],
['4','exam_instraction.php / start_exam.php','Instruction page or timed MCQ UI is displayed. start_exam.php reads configured duration.'],
['5','submitexam.php','Answers are checked against questions; user_answers and scores are stored; JSON returns score.'],
['6','pastexam.php / f_pdf.php','Student reviews results and can generate a PDF output.'],
], [0.55,2.1,4.55])

para(d,'Data model and retention','Heading 1')
para(d,'The physical schema is not checked into source control. The table model below is reconstructed from application queries and must be converted into a versioned SQL migration before the next release.')
table(d,['Table','Purpose','Key fields observed'],[
['registration','Student profile and credentials','registration_id, fullname, dob, email, phone, college, yearofstudy, passingoutyear, branch, skills, password'],
['exam','Assessment definition','exam_id, examtitle, examdescription, examdate, examtime, duration, gracestart, graceend, skillRequired'],
['sections','Exam sections and quotas','section_id, exam_id, section_name, noofquestions'],
['questions','MCQ question bank per exam/section','question_id, exam_id, section_id, question, option1–option4, correctanswer'],
['passkey','Per-student assessment access','passkey_id, registration_id, exam_id, passkey_n, is_used'],
['user_answers','Submitted response detail','registration_id, exam_id, question_id, selected_answer, correct_answer, is_correct'],
['scores','Aggregate submission result','registration_id, exam_id, score, total_questions, submitted_at'],
], [1.25,2.05,3.9])
callout(d,'Integrity recommendation','Define primary keys, unique keys (registration.email; scores.registration_id + scores.exam_id; passkey registration/exam/passkey), and foreign keys for every relationship. Take a database backup before running fix_db_schema.php or any schema operation.',GOLD)

para(d,'Setup, runbook and verification','Heading 1')
para(d,'Local setup (XAMPP)','Heading 2')
table(d,['Action','Expected result'],[
['1. Place project','Keep the code at C:\\xampp\\htdocs\\exam.'],
['2. Start services','Start Apache and MySQL from XAMPP Control Panel.'],
['3. Provision database','Create database exam and restore/create tables matching the data model above.'],
['4. Configure connection','Review config.php. Current development configuration is localhost / root / blank password / exam. Replace for non-local environments.'],
['5. Browse application','Open http://localhost/exam/login.php.'],
['6. Smoke test','Register a student, log in, create an exam, publish questions/passkeys, take an attempt, verify score/history/PDF.'],
], [2.0,5.2])
para(d,'Operational checks','Heading 2')
bullet(d,'Before an exam: verify the grace window, section question counts, passkeys and candidate eligibility; test one candidate journey end-to-end.')
bullet(d,'During an exam: monitor Apache/PHP error logs and MySQL availability; do not change question data or grace times without recording the impact.')
bullet(d,'After an exam: export results, back up the database, retain answers/scores per policy, and reset only with an approved data-retention decision.')
para(d,'Testing status','Heading 2')
para(d,'run-tests.ps1 attempts to execute tests\\run-tests.php using XAMPP PHP or a PHP binary on PATH. The expected test file was absent during repository inspection, so automated test coverage is currently not runnable. Establish smoke tests first, then add PHP unit/integration tests for authentication, exam windows, passkeys and scoring.')

para(d,'Security and support risks','Heading 1')
table(d,['Priority','Finding','Impact / required action'],[
['P0','Missing save_questions.php','Publishing questions/passkeys from add_question.php fails. Restore and test authenticated, validated endpoint.'],
['P0','Hardcoded admin credential in login.php','Credential exposure and no admin lifecycle. Move to database-backed roles; rotate secret immediately.'],
['P0','Admin route guards are incomplete/commented','Pages such as admin_student_scores.php can be reached without enforced authorization. Centralize role checks.'],
['P1','SQL interpolation in login.php and some IDs','Injection risk despite partial use of prepared statements. Convert all query inputs to prepared statements.'],
['P1','Credentials stored in config.php with root/blank password','Unsafe beyond local development. Use least-privilege DB user and environment configuration.'],
['P1','Client-generated, short passkeys','Predictable/low entropy and exposed in UI. Generate securely server-side; rate-limit verification.'],
['P1','No source-controlled schema/migrations','Environment drift and recovery risk. Create versioned SQL migrations and seed data.'],
['P2','Mixed legacy/duplicate exam pages','start_exam.php, startexam.php, test*.php and reference.php show inconsistent models. Decide a canonical path and retire dead code.'],
['P2','Debug log and schema fixer in web root','Potential data leakage or accidental schema changes. Remove from public access and restrict developer tools.'],
], [0.55,2.55,4.1])

para(d,'Recommended stabilization roadmap','Heading 1')
table(d,['Horizon','Deliverables','Exit criteria'],[
['0–2 days','Restore missing save_questions.php; protect all admin pages; rotate hardcoded admin credential; block public access to debug/log and schema tool.','Complete create → assign → attempt → score flow works with authorization.'],
['Week 1','Add schema SQL/migrations, backup/restore procedure, environment config, logging policy and manual smoke-test checklist.','A clean environment can be rebuilt and verified by a new engineer.'],
['Weeks 2–3','Refactor common bootstrap/auth/database helpers; replace interpolated SQL; centralize session and CSRF strategy.','Critical pages share consistent authentication and parameterized queries.'],
['Month 1','Automated test suite, passkey hardening, audit trail, role model, deployment playbook.','Release is repeatable, testable and supportable.'],
], [1.0,3.7,2.5])

para(d,'Ownership handover checklist','Heading 1')
table(d,['Check','Owner sign-off'],[
['Obtain current database export and validate restore into a non-production instance.',''],
['Confirm XAMPP/PHP/MySQL versions and document production-specific settings.',''],
['Restore or implement save_questions.php; test passkey assignment and persistence.',''],
['Replace hardcoded admin account and add centralized admin authorization.',''],
['Inventory/retire test*.php, reference.php, fix_db_schema.php and debug.log from web exposure.',''],
['Create schema migrations, backups, retention policy and incident contact path.',''],
['Execute the full admin and student smoke paths with representative data.',''],
['Publish a release checklist and test ownership for future changes.',''],
], [5.8,1.4])
para(d,'Appendix: source review notes','Heading 1')
bullet(d,'The application name appears as SmartExam Pro and ExamPortal in different pages; standardize branding as part of UI maintenance.')
bullet(d,'start_exam.php and startexam.php implement different exam experiences. The configured-duration implementation in start_exam.php is closer to the stated application intent; startexam.php includes a fixed 1-hour timer.')
bullet(d,'submitexam.php clears previous answers for the same user/exam but preserves the first scores row if it exists; define an explicit policy for resubmission/regrading.')
bullet(d,'upcoming_exam.php limits visibility to today and tomorrow and uses Asia/Kolkata timezone. Confirm that policy with exam operations.')
bullet(d,'The supplied context.md was useful background but contains inaccuracies relative to the source (for example, it references save_questions.php as if present). This KT document calls out those discrepancies.')
para(d,'Document end', 'Heading 2')
para(d,'This handover should be updated alongside the source whenever authentication, schema, examination policy, or deployment behavior changes.', color=MUTED, size=9.5)
d.save(OUT)
print(OUT)
